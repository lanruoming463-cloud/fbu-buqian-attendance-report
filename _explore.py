import openpyxl

files = [
    r"D:/Documents/Desktop/AI做每月人力数据/T+1确认及时性/劳务工-考勤记录-10.1-10.6-及时性.xlsx",
    r"D:/Documents/Desktop/AI做每月人力数据/T+1确认及时性/劳务工-考勤记录-10.7-及时性.xlsx",
    r"D:/Documents/Desktop/AI做每月人力数据/T+1确认及时性/正式工-考勤记录-10.1-10.6-及时性.xlsx",
    r"D:/Documents/Desktop/AI做每月人力数据/T+1确认及时性/正式工-考勤记录-10.7-及时性.xlsx",
]

for fp in files:
    print('=' * 90)
    print('FILE:', fp)
    try:
        wb = openpyxl.load_workbook(fp, read_only=True, data_only=True)
    except Exception as e:
        print('  LOAD ERR:', e)
        continue
    print('  SHEETS:', wb.sheetnames, '| max_row(est):', wb.active.max_row)
    for sn in wb.sheetnames:
        ws = wb[sn]
        print('  ---- SHEET:', repr(sn), 'dims:', ws.max_row, 'x', ws.max_column)
        n = 0
        for row in ws.iter_rows(values_only=True):
            if n < 3:
                vals = [('' if v is None else str(v)[:16]) for v in row[:70]]
                print('     r%d:' % n, vals)
            n += 1
            if n >= 3:
                break
    wb.close()
print('DONE')
