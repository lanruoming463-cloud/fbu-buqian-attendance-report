# -*- coding: utf-8 -*-
"""检查正式工-8月.xlsx：规范所需列、捷克区、行数与日期范围"""
import openpyxl, glob, sys
from collections import Counter

files = glob.glob(r"D:\Documents\Downloads\*正式工-8月*.xlsx")
print("files:", files)
wb = openpyxl.load_workbook(files[0], read_only=True)
print("sheets:", wb.sheetnames)
ws = wb[wb.sheetnames[0]]
rows = ws.iter_rows(values_only=True)
header = [str(c) if c is not None else "" for c in next(rows)]

# 找规范相关列
targets = ["考勤状态", "异常备注", "备注", "迟到", "早退", "考勤日期", " Czech", "捷克",
           "三级部门", "四级部门", "五级部门", "六级部门", "七级部门", "八级部门", "用工", "姓名", "工号"]
for i, h in enumerate(header):
    for t in targets:
        if t.strip() and t.strip().lower() in h.lower():
            print(f"col[{i}] = {h}")
            break

data = list(rows)
print("data rows:", len(data))

# 考勤日期列定位与范围
date_idx = None
for i, h in enumerate(header):
    if "考勤日期" in h:
        date_idx = i
        break
print("date_idx:", date_idx, "header:", header[date_idx] if date_idx is not None else None)
if date_idx is not None:
    days = Counter(str(r[date_idx])[:10] for r in data if r[date_idx] is not None)
    ks = sorted(days)
    print("date range:", ks[0], "->", ks[-1], " distinct days:", len(ks))

# 捷克检查：全表扫描各组织列
org_idx = [i for i, h in enumerate(header) if ("部门" in h or "组织" in h or "区域" in h or "大区" in h)]
print("org cols:", [(i, header[i]) for i in org_idx])
cz = Counter()
for r in data:
    for i in org_idx:
        v = r[i]
        if v and ("捷克" in str(v) or "czech" in str(v).lower() or "Czech" in str(v)):
            cz[header[i]] += 1
print("czech hits:", dict(cz))
